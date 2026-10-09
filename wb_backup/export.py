import hashlib
import json
import re
import zipfile
from pathlib import Path
from decimal import Decimal, InvalidOperation
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill, Alignment

LABELS = {"reportId": "ID отчёта", "rrdId": "ID операции", "nmId": "Артикул WB",
          "chrtId": "ID размера", "barcode": "Штрихкод", "warehouseId": "ID склада",
          "warehouseName": "Склад", "quantity": "Количество", "amount": "Остаток",
          "current": "Баланс", "for_withdraw": "Доступно к выводу", "currency": "Валюта",
          "date": "Дата", "dateFrom": "Начало периода", "dateTo": "Конец периода",
          "srid": "ID заказа", "vendorCode": "Артикул продавца", "retailAmount": "Сумма продажи",
          "inWayToClient": "В пути к покупателю", "inWayFromClient": "В пути от покупателя"}


def safe_name(value):
    return re.sub(r'[^\w. -]', '_', str(value), flags=re.UNICODE).strip('. ')[:100] or 'file'


def write_table(path: Path, records, metadata, max_rows=50000):
    """One source page per workbook; streaming sheets with safe literal strings."""
    path.parent.mkdir(parents=True, exist_ok=True)
    columns = list(dict.fromkeys(key for row in records for key in row))
    if len(columns) > 16384:
        raise ValueError("Слишком много столбцов для Excel; исходные данные сохранены")
    workbook = Workbook(write_only=True)
    summary = workbook.create_sheet("Описание")
    summary.column_dimensions['A'].width = 28
    summary.column_dimensions['B'].width = 110
    for key, value in metadata.items():
        append_literal(summary, [key, value])
    for index in range(1, len(metadata) + 3):
        summary.row_dimensions[index].height = 42
    append_literal(summary, ["Строк", len(records)])
    append_literal(summary, ["Важно", "Часть выгрузки; полнота и периоды указаны в реестре архива. Баланс и остатки — снимки, а не история."])
    sheet = None
    for index, record in enumerate(records):
        if index % max_rows == 0:
            sheet = workbook.create_sheet(f"Данные {index // max_rows + 1}")
            sheet.freeze_panes = "A2"
            for col in range(1, len(columns) + 1):
                from openpyxl.utils import get_column_letter
                sheet.column_dimensions[get_column_letter(col)].width = 24
            header = []
            for key in columns:
                cell = WriteOnlyCell(sheet, value=LABELS.get(key, key) + (f" [{key}]" if key in LABELS else ""))
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="16324F")
                cell.alignment = Alignment(wrap_text=True)
                header.append(cell)
            sheet.append(header)
            sheet.row_dimensions[1].height = 34
        values = []
        for key in columns:
            value = record.get(key)
            identifier = key.lower().endswith("id") or key.lower() in {"barcode", "skus", "inn", "tin"}
            if isinstance(value, (dict, list)):
                value = json.dumps(value, ensure_ascii=False)
            if value is not None and (identifier or isinstance(value, int) and abs(value) >= 10 ** 15):
                value = str(value)
            elif isinstance(value, str) and any(word in key.lower() for word in
                    ('amount', 'sum', 'fee', 'price', 'payment', 'deliveryrub', 'forpay', 'penalty', 'commission', 'retail', 'storage')):
                try:
                    number = Decimal(value)
                    # Excel cannot preserve more than 15 significant decimal digits.
                    if number.is_finite() and len(number.as_tuple().digits) <= 15:
                        value = number
                except InvalidOperation:
                    pass
            values.append(value)
        append_literal(sheet, values)
    if not records:
        append_literal(workbook.create_sheet("Данные"), ["За запрошенный период API не вернул строк"])
    temp = path.with_suffix(".tmp.xlsx")
    workbook.save(temp)
    temp.replace(path)


def append_literal(sheet, values):
    cells = []
    for value in values:
        cell = WriteOnlyCell(sheet, value=value)
        cell.alignment = Alignment(vertical='top', wrap_text=sheet.title == 'Описание')
        if isinstance(value, str):
            # Keep exact raw data separately; Excel has a 32767-character cell limit.
            cell.value = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', value)[:32767]
            cell.data_type = "s"  # no formula execution from external product text
        cells.append(cell)
    sheet.append(cells)


def package(root: Path, output: Path, maximum: int):
    """Independent ZIP parts. Oversized individual files are byte chunks with a join guide."""
    output.mkdir(parents=True, exist_ok=True)
    files = [p for p in sorted(root.rglob('*')) if p.is_file() and '.tmp' not in p.name and p.name != 'checksums.json']
    manifest = []
    for p in files:
        digest = hashlib.sha256()
        with p.open('rb') as stream:
            for chunk in iter(lambda: stream.read(1024 * 1024), b''):
                digest.update(chunk)
        manifest.append({"file": p.relative_to(root).as_posix(), "bytes": p.stat().st_size, "sha256": digest.hexdigest()})
    from .storage import atomic_json
    atomic_json(root / 'checksums.json', manifest)
    files.append(root / 'checksums.json')
    # Stored ZIP guarantees a size bound even for already-compressed documents.
    budget = maximum - 1024 * 1024
    entries = []
    for p in files:
        if p.stat().st_size <= budget:
            entries.append((p.relative_to(root).as_posix(), p, None))
        else:
            with p.open('rb') as stream:
                part = 0
                while block := stream.read(budget // 2):
                    part += 1
                    chunk = output / f'chunk_{len(entries):08d}'
                    chunk.write_bytes(block)
                    entries.append((p.relative_to(root).as_posix() + f'.chunk{part:05d}', chunk, chunk))
    guide = ('Распакуйте все ZIP в одну папку. Файлы *.chunk00001, *.chunk00002 и далее — части одного файла.\n'
             'Для восстановления: python join_chunks.py <папка>. Сверьте checksums.json.\n')
    archives = []
    current = None
    size = 0
    for name, p, disposable in entries:
        overhead = len(name.encode('utf-8')) * 2 + 200
        if current is None or size + p.stat().st_size + overhead > budget:
            if current:
                current.close()
            path = output / f'wb_archive_{len(archives) + 1:03d}.zip'
            archives.append(path)
            current = zipfile.ZipFile(path, 'w', compression=zipfile.ZIP_STORED)
            current.writestr('ПРОЧИТАЙТЕ.txt', guide)
            current.write(Path(__file__).with_name('join_chunks.py'), 'join_chunks.py')
            size = 10000
        current.write(p, name)
        size += p.stat().st_size + overhead
        if disposable:
            p.unlink()
    if current:
        current.close()
    if any(p.stat().st_size > maximum for p in archives):
        raise ValueError('Архив превысил лимит; файлы сохранены локально')
    return archives
